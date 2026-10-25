"""Recalcule chaque ligne de la feuille Results du classeur et la compare aux valeurs écrites par le VBA.

Usage : python3 scripts/validate_workbook.py chemin/vers/classeur.xlsm
Nécessite openpyxl (pip install openpyxl).
"""

import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from winter_challenge.scoring import COUNT_LABELS, game_handicap, hole_counts, score_round  # noqa: E402

# Colonnes (indices 0) de la feuille Results
HOLES = slice(4, 22)  # E..V
EXPECTED = {
    "net_total": 23,  # X
    "stableford_gross": 24,  # Y
    "stableford_net": 25,  # Z
    "points_gross": 26,  # AA
    "points_net": 27,  # AB
    "hcp_after": 29,  # AD
}
HCP_BEFORE = 28  # AC


def load_courses(wb):
    """Nom du parcours -> (pars, strokes) sur 18 trous."""
    courses = {}
    for row in wb["Courses"].iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        pars = list(row[2:11]) + list(row[12:21])  # C..K, M..U
        strokes = list(row[23:41])  # X..AO
        courses[row[1]] = (pars, strokes)
    return courses


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    courses = load_courses(wb)
    total = 0
    diffs = []
    for r, row in enumerate(wb["Results"].iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        total += 1
        if row[3] not in courses:
            diffs.append((r, row[1], row[2], "parcours inconnu: %s" % row[3]))
            continue
        pars, strokes = courses[row[3]]
        scores = [int(v) if v not in (None, "") else 0 for v in row[HOLES]]
        res = score_round(scores, pars, strokes, float(row[HCP_BEFORE]))
        for i, label in enumerate(COUNT_LABELS):
            got, want = hole_counts(scores, pars)[label], row[30 + i]  # AE..AL
            if want is None or got != want:
                diffs.append((r, row[1], row[2], f"{label}: calculé {got}, classeur {want}"))
        for field, col in EXPECTED.items():
            got, want = getattr(res, field), row[col]
            if want is None or abs(float(got) - float(want)) > 1e-9:
                diffs.append((r, row[1], row[2], f"{field}: calculé {got}, classeur {want}"))
    players = 0
    for r, row in enumerate(wb["Players"].iter_rows(min_row=2, values_only=True), start=2):
        try:
            real, want = float(row[6]), row[7]
        except (TypeError, ValueError):
            continue  # handicap pas encore connu (« ? »)
        if real == 0:
            print(f"  note : {row[2]} a un handicap réel de 0, cas limite saisi à la main dans le classeur (la règle donne -1) : ignoré")
            continue
        players += 1
        if game_handicap(real) != want:
            diffs.append((r, row[2], "Players", f"handicap de jeu: calculé {game_handicap(real)}, classeur {want}"))
    print(f"{total} résultats et {players} handicaps de départ vérifiés, {len(diffs)} écart(s)")
    for d in diffs[:40]:
        print("  ligne %s  %s  %s  %s" % d)
    return 1 if diffs else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
