"""Écriture du classeur Excel (.xlsx, sans macros) à partir d'une saison calculée."""

from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .scoring import COUNT_LABELS
from .season import STATUS_RETAINED, Season

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
COUNTED_FILL = PatternFill("solid", fgColor="DFF0D8")
REPLACED_FILL = PatternFill("solid", fgColor="F4B6B6")
RETAINED_FILL = PatternFill("solid", fgColor="DFF0D8")
GREY_FONT = Font(color="808080")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

FR_COUNTS = ["Albatros", "Eagles", "Birdies", "Pars", "Bogeys", "Doubles", "Triples", "Plus"]
STATUS_FR_WEEK = {"Completed": "Terminée", "In Progress": "En cours"}


def style_header(ws, row: int, first_col: int, last_col: int):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row, c)
        cell.fill, cell.font, cell.alignment, cell.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
    ws.row_dimensions[row].height = 30


def set_widths(ws, widths: dict):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def evolution_text(value) -> str:
    if value is None:
        return ""
    if value == "Nouveau":
        return "Nouveau"
    return f"▲ {value}" if value > 0 else (f"▼ {-value}" if value < 0 else "=")


def write_ranking(wb: Workbook, season: Season, kind: str, title: str, season_label: str):
    ws = wb.create_sheet(title)
    week_numbers = sorted({w.number for w in season.weeks}, reverse=True)  # la plus récente à gauche, comme avant
    first_week_col = 8
    row = 1
    for category in dict.fromkeys(p.category for p in season.players):
        rows = season.rankings[(category, kind)]
        ws.cell(row, 1, f"Classement {'net Stableford' if kind == 'net' else 'brut'} {season_label} - {category}").font = TITLE_FONT
        row += 1
        headers = ["Rang", "Évol.", "Joueur", "Hdc départ", "Hdc actuel", "Nb semaines",
                   f"Total ({season.best_n} meilleurs)"] + [f"WC{n}" for n in week_numbers]
        for c, h in enumerate(headers, start=1):
            ws.cell(row, c, h)
        style_header(ws, row, 1, len(headers))
        row += 1
        for r in rows:
            values = [r.rank, evolution_text(r.evolution), r.player, r.start_hcp, r.current_hcp, r.nb_weeks, r.total]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row, c, v)
                cell.border, cell.alignment = BORDER, CENTER if c != 3 else Alignment(horizontal="left")
            ws.cell(row, 7).font = Font(bold=True)
            for i, n in enumerate(week_numbers):
                cell = ws.cell(row, first_week_col + i, r.points.get(n))
                cell.border, cell.alignment = BORDER, CENTER
                if n in r.counted_weeks:
                    cell.fill = COUNTED_FILL
                if n == r.replaced_week:
                    cell.fill = REPLACED_FILL
            row += 1
        row += 2
    set_widths(ws, {"A": 7, "B": 10, "C": 20, "D": 11, "E": 11, "F": 12, "G": 18})
    for i in range(len(week_numbers)):
        ws.column_dimensions[get_column_letter(first_week_col + i)].width = 7
    ws.freeze_panes = "D3"


def write_results(wb: Workbook, season: Season):
    ws = wb.create_sheet("Résultats")
    headers = (["Semaine", "Catégorie", "Joueur", "N° tentative", "Parcours", "Départ", "Trous"]
               + [f"T{i}" for i in range(1, 19)]
               + ["Brut", "Net", "Stableford brut", "Stableford net", "Points bruts", "Points nets",
                  "Hdc avant", "Hdc après"] + FR_COUNTS + ["Statut"])
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for a in sorted(season.attempts, key=lambda a: (a.week, a.category, a.player, a.number)):
        res = a.result
        row = [f"WC{a.week}", a.category, a.player, a.number, a.course, a.tee, a.section] + [s or None for s in a.scores]
        if res:
            row += [res.gross_total, res.net_total, res.stableford_gross, res.stableford_net, res.points_gross,
                    res.points_net, res.hcp_before, res.hcp_after]
        else:
            row += [sum(a.scores)] + [None] * 7
        row += [a.counts[label] for label in COUNT_LABELS] + [a.status]
        ws.append(row)
        if a.status == STATUS_RETAINED:
            for c in range(1, len(headers) + 1):
                ws.cell(ws.max_row, c).fill = RETAINED_FILL
        elif a.status != "Non retenu":
            for c in range(1, len(headers) + 1):
                ws.cell(ws.max_row, c).font = GREY_FONT
    set_widths(ws, {"A": 9, "B": 12, "C": 18, "D": 11, "E": 28, "F": 10, "G": 8})
    for i in range(18):
        ws.column_dimensions[get_column_letter(8 + i)].width = 4.5
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions


def write_weeks(wb: Workbook, season: Season):
    ws = wb.create_sheet("Semaines")
    headers = ["Semaine", "Tour", "Tournoi SGT", "Nom SGT", "Parcours", "Début", "Fin", "Statut", "Joueurs", "Tentatives",
               "Meilleur Stableford brut", "Joueur(s)", "Meilleur Stableford net", "Joueur(s)"] + FR_COUNTS
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for w in sorted(season.weeks, key=lambda w: (w.number, w.tour_id)):
        attempts = [a for a in season.attempts if a.week == w.number and a.category == w.category]
        retained = [a for (cat, _, wk), a in season.best.items() if wk == w.number and cat == w.category]
        row = [f"WC{w.number}", w.category, w.tournament_id, w.name, w.course, w.start, w.end,
               STATUS_FR_WEEK.get(w.status, w.status), len({a.player for a in attempts}), len(attempts)]
        for kind in ("stableford_gross", "stableford_net"):
            best = max((getattr(a.result, kind) for a in retained), default=None)
            who = ", ".join(sorted(a.player for a in retained if getattr(a.result, kind) == best)) if retained else ""
            row += [best, who]
        row += [sum(a.counts[label] for a in retained) for label in COUNT_LABELS]
        ws.append(row)
    set_widths(ws, {"A": 9, "B": 12, "C": 12, "D": 14, "E": 28, "F": 12, "G": 12, "H": 10, "K": 14, "L": 24, "M": 14, "N": 24})
    ws.freeze_panes = "B2"


def write_players(wb: Workbook, season: Season):
    ws = wb.create_sheet("Joueurs")
    headers = ["Joueur", "Catégorie", "Handicap réel", "Hdc de jeu départ", "Hdc de jeu actuel", "Source du handicap"]
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for p in sorted(season.players, key=lambda p: (p.category, p.name.lower())):
        ws.append([p.name, p.category, p.real_hcp, p.start_hcp, p.current_hcp, p.hcp_source])
    set_widths(ws, {"A": 20, "B": 14, "C": 14, "D": 16, "E": 16, "F": 26})
    ws.freeze_panes = "A2"


def write_params(wb: Workbook, season: Season, season_label: str, tours: list):
    ws = wb.create_sheet("Paramètres")
    rows = [
        ("Saison", season_label),
        ("Tours SGT utilisés", ", ".join(f"{t['tour_id']} ({t['name']} → {t['label']})" for t in tours)),
        ("Résultats comptés au classement", season.best_n),
        ("Tentatives par semaine", season.max_attempts),
        ("Fichier généré le", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("", ""),
        ("Comment lire les classements", ""),
        ("Vert", "résultat compté parmi les meilleurs"),
        ("Rouge", "plus petit résultat compté : le premier à être remplacé (le plus ancien à égalité)"),
        ("Points", "100 + (Stableford − 36) × 10 sur 18 trous, − 18 sur 9 trous, minimum 30"),
        ("Tentative retenue", "meilleur Stableford net de la semaine (à égalité, la plus récente)"),
        ("Handicap de jeu", "+1 sous 100 points nets, sinon − (points − 100) / 10, mis à jour à la fin de chaque semaine"),
        ("Handicap de départ", "50 % du handicap réel (fichier joueurs, sinon handicap fixé par le club dans SGT)"),
        ("", ""),
        ("Ce fichier est généré par Python à partir de l'API SGT : ne pas le corriger à la main.", ""),
    ]
    for r in rows:
        ws.append(list(r))
    ws["A7"].font = Font(bold=True)
    ws["A15"].font = Font(italic=True)
    ws["B8"].fill, ws["B9"].fill = COUNTED_FILL, REPLACED_FILL
    set_widths(ws, {"A": 34, "B": 100})


def write_workbook(season: Season, tours: list, season_label: str, path: str):
    wb = Workbook()
    wb.remove(wb.active)
    write_ranking(wb, season, "net", "Classement Net", season_label)
    write_ranking(wb, season, "brut", "Classement Brut", season_label)
    write_results(wb, season)
    write_weeks(wb, season)
    write_players(wb, season)
    write_params(wb, season, season_label, tours)
    wb.save(path)
